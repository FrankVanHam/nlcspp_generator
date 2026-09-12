from __future__ import annotations

from collections import OrderedDict
from pathlib import Path
from typing import Any

import openpyxl

from .model import Feature, Property


SHEET_NAME = "Informatiemodel"
HEADER_ROW = 2
REQUIRED_HEADERS = {
    "Tabel",
    "Attribuut",
    "DATATYPE",
    "LENGTH",
    "Keuzelijst",
    "Verplicht XSD",
    "Geometrietype",
}
SCALAR_TYPES = {
    "DATE ONLY": "xs:date",
    "DATE": "xs:dateTime",
    "DOUBLE": "xs:decimal",
    "FLOAT": "xs:decimal",
    "LONG": "xs:integer",
}
GEOMETRY_TYPES = {
    "puntgeometrie": "gml:PointPropertyType",
    "lijngeometrie": "gml:CurvePropertyType",
    "vlakgeometrie": "gml:SurfacePropertyType",
}
IDREF_PROPERTIES = {"MantelbuisID", "InhoudID", "AssetObjectID"}

# The original schema's lexical order is part of its public byte-level contract.
# New workbook object types are appended in their source order.
LEGACY_FEATURE_ORDER = (
    "LSkast", "LSmof", "MSstation", "MSoverdrachtspunt", "LSkabel", "MSkabel",
    "OVLoverdrachtspunt", "LSoverdrachtspunt", "MSmof", "Amantelbuis",
    "AmantelbuisInhoud", "Amarkeringsobject", "KBgelijkrichter", "KBmeetpaal",
    "KBmeetdraad", "KBanode", "KBdrainage", "KBobject", "Tkabel", "Tmof",
    "Toverdrachtspunt", "Tstation", "Tbuis", "HSkabel", "HSmof", "HSstation",
    "Eoliedrukleiding", "Eoliedrukinstallatie", "TbuisAppendage", "Eaardmof",
    "Eaarddraad", "Eaardpen", "AbeschermingVlak", "Aopmerking", "AbestandBijlage",
    "Amaaiveldhoogte", "Akunstwerk", "Aaanlegtechniek", "AinUittredepunt",
    "Averplaatsing", "Gstation", "Gleiding", "Gafsluiter", "GtStuk",
    "Goverdrachtspunt", "Gaftakzadel", "Geindstuk", "GappendageOverig",
    "Gisolatiestuk", "Govergangsstuk", "Gmeetpunt", "Gsifon", "GverticaleBocht",
    "Gontspanningselement", "Gblaasgatzadel", "Gleidingafblaas",
)
LEGACY_SELECTION_FEATURE_ORDER = (
    "LSkast", "LSmof", "LSkabel", "MSkabel", "OVLoverdrachtspunt",
    "LSoverdrachtspunt", "MSmof", "Amantelbuis", "Gstation", "Gleiding",
    "Gafsluiter", "GtStuk", "Goverdrachtspunt", "Gaftakzadel", "Geindstuk",
    "GappendageOverig", "Gisolatiestuk", "Govergangsstuk", "Gmeetpunt",
    "Gontspanningselement", "Gleidingafblaas", "KBgelijkrichter", "KBmeetpaal",
    "KBmeetdraad", "KBanode", "KBdrainage", "KBobject", "Tbuis", "HSkabel",
    "HSmof", "TbuisAppendage", "Eaarddraad", "Eaardpen", "AbeschermingVlak",
    "Akunstwerk", "Aaanlegtechniek",
)


def load_features(path: Path) -> tuple[Feature, ...]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise ValueError(f"Missing sheet {SHEET_NAME!r} in {path.name}")
        sheet = workbook[SHEET_NAME]
        rows = sheet.iter_rows(values_only=True)
        next(rows)
        header_values = next(rows)
        headers = {
            str(value): column_index
            for column_index, value in enumerate(header_values)
            if value is not None
        }
        missing = sorted(REQUIRED_HEADERS - headers.keys())
        if missing:
            raise ValueError(f"Missing columns in sheet {SHEET_NAME!r}: {', '.join(missing)}")

        properties_by_feature: OrderedDict[str, list[Property]] = OrderedDict()
        geometry_by_feature: dict[str, Property] = {}
        previous_geometry_by_feature: dict[str, str] = {}
        for row_number, row in enumerate(rows, HEADER_ROW + 1):
            feature_name = _text(row[headers["Tabel"]])
            property_name = _text(row[headers["Attribuut"]])
            if not feature_name or not property_name:
                continue
            geometry = _text(row[headers["Geometrietype"]])
            property_ = _read_property(
                row,
                headers,
                row_number,
                feature_name,
                property_name,
                geometry or previous_geometry_by_feature.get(feature_name, ""),
            )
            if geometry:
                previous_geometry_by_feature[feature_name] = geometry
            properties = properties_by_feature.setdefault(feature_name, [])
            if property_.type_name.startswith("gml:"):
                if feature_name in geometry_by_feature:
                    raise ValueError(f"Multiple geometry rows for {feature_name!r}")
                geometry_by_feature[feature_name] = property_
            else:
                properties.append(property_)

        return tuple(
            Feature(
                name=feature_name,
                properties=(
                    Property("Handle", "xs:string", required=False),
                    *properties,
                    geometry_by_feature[feature_name],
                ),
            )
            for feature_name, properties in properties_by_feature.items()
        )
    finally:
        workbook.close()


def domain_order(features: tuple[Feature, ...]) -> tuple[str, ...]:
    seen: set[str] = set()
    ordered = []
    for feature in features:
        for property_ in feature.properties:
            if ":" not in property_.type_name and property_.type_name not in seen:
                ordered.append(property_.type_name)
                seen.add(property_.type_name)
    return tuple(ordered)


def base_schema_domain_order(
    project_feature: Feature, features: tuple[Feature, ...]
) -> tuple[str, ...]:
    kb_index = next(
        (index for index, feature in enumerate(features) if feature.name.startswith("KB")),
        len(features),
    )
    shared_gas_domain = "BewerkingGas"
    ordered = list(domain_order(features[:kb_index]))
    ordered = [name for name in ordered if name != shared_gas_domain]
    for name in domain_order((project_feature,)):
        if name not in ordered:
            ordered.append(name)
    if any(
        property_.type_name == shared_gas_domain
        for feature in features
        for property_ in feature.properties
    ):
        ordered.append(shared_gas_domain)
    for name in domain_order(features[kb_index:]):
        if name not in ordered:
            ordered.append(name)
    return tuple(ordered)


def selection_schema_domain_order(features: tuple[Feature, ...]) -> tuple[str, ...]:
    by_name = {feature.name: feature for feature in features}
    ordered_features = [
        by_name[name] for name in LEGACY_SELECTION_FEATURE_ORDER if name in by_name
    ]
    ordered_names = {feature.name for feature in ordered_features}
    ordered_features.extend(
        feature for feature in features if feature.name not in ordered_names
    )
    return domain_order(tuple(ordered_features))


def order_features(features: tuple[Feature, ...]) -> tuple[Feature, ...]:
    by_name = {feature.name: feature for feature in features}
    ordered = [by_name[name] for name in LEGACY_FEATURE_ORDER if name in by_name]
    ordered_names = {feature.name for feature in ordered}
    ordered.extend(feature for feature in features if feature.name not in ordered_names)
    return tuple(ordered)


def _read_property(
    row: tuple[Any, ...],
    headers: dict[str, int],
    row_number: int,
    feature_name: str,
    property_name: str,
    geometry: str,
) -> Property:
    data_type = _text(row[headers["DATATYPE"]])
    required_value = _text(row[headers["Verplicht XSD"]])
    if required_value not in {"NULLABLE", "NON_NULLABLE"}:
        raise ValueError(
            f"Unsupported requiredness {required_value!r} in {SHEET_NAME}!{row_number}"
        )
    required = required_value == "NON_NULLABLE"

    if data_type == "SHAPE":
        try:
            type_name = GEOMETRY_TYPES[geometry]
        except KeyError as error:
            raise ValueError(
                f"Unsupported geometry {geometry!r} in {SHEET_NAME}!{row_number}"
            ) from error
        return Property(property_name, type_name, required, source_row=row_number)

    domain_name = _text(row[headers["Keuzelijst"]])
    if property_name == "ID" and feature_name != "GtStuk":
        type_name = "xs:ID"
        max_length = None
    elif property_name in IDREF_PROPERTIES:
        type_name = "xs:IDREF"
        max_length = None
    elif domain_name and domain_name != "nvt":
        type_name = domain_name
        max_length = None
    elif data_type == "TEXT":
        type_name = "xs:string"
        max_length = _integer(row[headers["LENGTH"]], row_number)
    else:
        try:
            type_name = SCALAR_TYPES[data_type]
        except KeyError as error:
            raise ValueError(
                f"Unsupported datatype {data_type!r} in {SHEET_NAME}!{row_number}"
            ) from error
        max_length = None
    return Property(property_name, type_name, required, max_length, row_number)


def _integer(value: Any, row_number: int) -> int:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(f"Expected numeric LENGTH in {SHEET_NAME}!{row_number}")
    if isinstance(value, float) and not value.is_integer():
        raise ValueError(f"Expected integral LENGTH in {SHEET_NAME}!{row_number}")
    return int(value)


def _text(value: Any) -> str:
    return "" if value is None else str(value)