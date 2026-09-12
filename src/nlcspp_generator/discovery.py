from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re
from typing import Iterator

import openpyxl
import xlrd


SUPPORTED_SUFFIXES = frozenset({".xls", ".xlsx"})


@dataclass(frozen=True)
class SheetInventory:
    name: str
    row_count: int
    column_count: int


@dataclass(frozen=True)
class WorkbookInventory:
    path: Path
    sheets: tuple[SheetInventory, ...]


@dataclass(frozen=True)
class InputSources:
    information_model: Path
    domain_catalog: Path
    version: str


def discover_workbooks(input_directory: Path) -> tuple[Path, ...]:
    if not input_directory.is_dir():
        raise NotADirectoryError(f"Input directory does not exist: {input_directory}")

    return tuple(
        path
        for path in sorted(input_directory.iterdir(), key=lambda candidate: candidate.name.casefold())
        if path.is_file()
        and not path.name.startswith("~$")
        and path.suffix.casefold() in SUPPORTED_SUFFIXES
    )


def inventory_workbook(path: Path) -> WorkbookInventory:
    suffix = path.suffix.casefold()
    if suffix == ".xlsx":
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
        try:
            sheets = tuple(
                SheetInventory(sheet.title, sheet.max_row, sheet.max_column)
                for sheet in workbook.worksheets
            )
        finally:
            workbook.close()
    elif suffix == ".xls":
        workbook = xlrd.open_workbook(path, on_demand=True)
        try:
            sheets = tuple(
                SheetInventory(sheet.name, sheet.nrows, sheet.ncols)
                for sheet in _iter_xls_sheets(workbook)
            )
        finally:
            workbook.release_resources()
    else:
        raise ValueError(f"Unsupported workbook format: {path.suffix}")

    return WorkbookInventory(path=path, sheets=sheets)


def inventory_directory(input_directory: Path) -> tuple[WorkbookInventory, ...]:
    return tuple(inventory_workbook(path) for path in discover_workbooks(input_directory))


def locate_input_sources(input_directory: Path) -> InputSources:
    inventories = inventory_directory(input_directory)
    information_models = [
        inventory.path
        for inventory in inventories
        if {sheet.name for sheet in inventory.sheets} >= {"Informatiemodel", "Overzicht Objecten"}
    ]
    domain_catalogs = [
        inventory.path
        for inventory in inventories
        if {sheet.name for sheet in inventory.sheets}
        >= {"domains", "domain_values", "domain_associations"}
    ]
    information_model = _exactly_one("information model", information_models)
    domain_catalog = _exactly_one("domain catalog", domain_catalogs)
    version_match = re.search(r"v(\d+(?:\.\d+)*)", information_model.name, re.IGNORECASE)
    if version_match is None:
        raise ValueError(f"Cannot infer schema version from {information_model.name!r}")
    return InputSources(information_model, domain_catalog, version_match.group(1))


def _iter_xls_sheets(workbook: xlrd.book.Book) -> Iterator[xlrd.sheet.Sheet]:
    for sheet_name in workbook.sheet_names():
        yield workbook.sheet_by_name(sheet_name)


def _exactly_one(role: str, paths: list[Path]) -> Path:
    if len(paths) != 1:
        found = ", ".join(path.name for path in paths) or "none"
        raise ValueError(f"Expected exactly one {role} workbook; found {found}")
    return paths[0]