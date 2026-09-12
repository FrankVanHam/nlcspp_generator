from __future__ import annotations

import argparse
from pathlib import Path

import openpyxl
import xlrd

from .discovery import discover_workbooks


def print_inventory(input_directory: Path, row_limit: int = 30) -> None:
    for path in discover_workbooks(input_directory):
        print(f"\n=== {path.name} ===")
        if path.suffix.casefold() == ".xlsx":
            _print_xlsx(path, row_limit)
        else:
            _print_xls(path, row_limit)


def _print_xlsx(path: Path, row_limit: int) -> None:
    workbook = openpyxl.load_workbook(path, read_only=False, data_only=False)
    try:
        for sheet in workbook.worksheets:
            print(
                f"-- {sheet.title!r}: rows={sheet.max_row}, "
                f"columns={sheet.max_column}, merged={list(sheet.merged_cells.ranges)}"
            )
            shown = 0
            for row in sheet.iter_rows():
                populated = [
                    f"{cell.coordinate}={cell.value!r}"
                    for cell in row
                    if cell.value is not None
                ]
                if populated:
                    print(" | ".join(populated))
                    shown += 1
                    if shown >= row_limit:
                        break
            comments = [
                (cell.coordinate, cell.comment.text)
                for row in sheet.iter_rows()
                for cell in row
                if cell.comment is not None
            ]
            if comments:
                print(f"comments={comments!r}")
    finally:
        workbook.close()


def _print_xls(path: Path, row_limit: int) -> None:
    workbook = xlrd.open_workbook(path, on_demand=True)
    try:
        for sheet_name in workbook.sheet_names():
            sheet = workbook.sheet_by_name(sheet_name)
            print(f"-- {sheet.name!r}: rows={sheet.nrows}, columns={sheet.ncols}")
            shown = 0
            for row_index in range(sheet.nrows):
                populated = [
                    f"{xlrd.formula.cellname(row_index, column_index)}="
                    f"{sheet.cell_value(row_index, column_index)!r}"
                    for column_index in range(sheet.ncols)
                    if sheet.cell_type(row_index, column_index)
                    not in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK)
                ]
                if populated:
                    print(" | ".join(populated))
                    shown += 1
                    if shown >= row_limit:
                        break
    finally:
        workbook.release_resources()


def main() -> None:
    parser = argparse.ArgumentParser(description="Inspect NLCS Excel workbook structure")
    parser.add_argument("input_directory", type=Path)
    parser.add_argument("--rows", type=int, default=30)
    arguments = parser.parse_args()
    print_inventory(arguments.input_directory, arguments.rows)


if __name__ == "__main__":
    main()