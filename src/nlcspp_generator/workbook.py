from __future__ import annotations

import re
from collections import defaultdict
from pathlib import Path

import xlrd
from openpyxl import load_workbook

from .model import AttributeDefinition, DomainDefinition, FeatureDefinition, InformationModel


MODEL_ORDER = (
    "LSkast", "LSmof", "MSstation", "MSoverdrachtspunt", "LSkabel", "MSkabel",
    "OVLoverdrachtspunt", "LSoverdrachtspunt", "MSmof", "Amantelbuis",
    "AmantelbuisInhoud", "AprojectReferentie", "Amarkeringsobject", "KBgelijkrichter",
    "KBmeetpaal", "KBmeetdraad", "KBanode", "KBdrainage", "KBobject", "Tkabel",
    "Tmof", "Toverdrachtspunt", "Tstation", "Tbuis", "HSkabel", "HSmof", "HSstation",
    "Eoliedrukleiding", "Eoliedrukinstallatie", "TbuisAppendage", "Eaardmof", "Eaarddraad",
    "Eaardpen", "AbeschermingVlak", "Aopmerking", "AbestandBijlage", "Amaaiveldhoogte",
    "Akunstwerk", "Aaanlegtechniek", "AinUittredepunt", "Averplaatsing", "Gstation",
    "Gleiding", "Gafsluiter", "GtStuk", "Goverdrachtspunt", "Gaftakzadel", "Geindstuk",
    "GappendageOverig", "Gisolatiestuk", "Govergangsstuk", "Gmeetpunt", "Gsifon",
    "GverticaleBocht", "Gontspanningselement", "Gblaasgatzadel", "Gleidingafblaas",
)

EXPECTED_HEADERS = {
    "Tabel", "Attribuut", "DATATYPE", "LENGTH", "Keuzelijst", "Verplicht XSD",
    "Geometrietype",
}


class WorkbookError(ValueError):
    pass


def _version_from_name(path: Path) -> str:
    match = re.search(r"v(\d+(?:\.\d+)*)", path.stem, re.IGNORECASE)
    if not match:
        raise WorkbookError(f"Cannot infer model version from {path.name}; pass --version")
    return match.group(1)


def _read_features(path: Path) -> tuple[FeatureDefinition, ...]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "Informatiemodel" not in workbook.sheetnames:
        raise WorkbookError(f"Missing worksheet 'Informatiemodel' in {path}")
    sheet = workbook["Informatiemodel"]
    headers = {cell.value: cell.column - 1 for cell in sheet[2] if cell.value}
    missing = EXPECTED_HEADERS - headers.keys()
    if missing:
        raise WorkbookError(f"Missing Informatiemodel columns: {', '.join(sorted(missing))}")

    by_feature: dict[str, list[AttributeDefinition]] = defaultdict(list)
    seen: set[tuple[str, str]] = set()
    for row_number, row in enumerate(sheet.iter_rows(min_row=3, values_only=True), 3):
        feature = row[headers["Tabel"]]
        name = row[headers["Attribuut"]]
        nullable = row[headers["Verplicht XSD"]]
        if not feature or not name or not nullable:
            continue
        key = (str(feature), str(name))
        if key in seen:
            raise WorkbookError(f"Duplicate XSD attribute {feature}.{name} at row {row_number}")
        seen.add(key)
        raw_length = row[headers["LENGTH"]]
        domain = row[headers["Keuzelijst"]]
        by_feature[str(feature)].append(AttributeDefinition(
            feature_name=str(feature),
            name=str(name),
            data_type=str(row[headers["DATATYPE"]]).strip().upper(),
            length=int(raw_length) if raw_length not in (None, "") else None,
            domain_name=None if domain in (None, "", "nvt") else str(domain),
            required=str(nullable).strip().upper() == "NON_NULLABLE",
            geometry_type=(str(row[headers["Geometrietype"]])
                           if row[headers["Geometrietype"]] else None),
            source_row=row_number,
        ))

    unknown = [name for name in by_feature if name not in MODEL_ORDER]
    for name, attributes in by_feature.items():
        by_feature[name] = [attribute for attribute in attributes if attribute.name != "Geometry"] + [
            attribute for attribute in attributes if attribute.name == "Geometry"
        ]
    ordered_names = [name for name in MODEL_ORDER if name in by_feature] + unknown
    missing_features = [name for name in MODEL_ORDER if name not in by_feature]
    if missing_features:
        raise WorkbookError(f"Missing expected XSD objects: {', '.join(missing_features)}")
    return tuple(FeatureDefinition(name, tuple(by_feature[name])) for name in ordered_names)


def _read_domains(path: Path) -> dict[str, DomainDefinition]:
    workbook = xlrd.open_workbook(path)
    try:
        definitions = workbook.sheet_by_name("domains")
        values_sheet = workbook.sheet_by_name("domain_values")
    except xlrd.biffh.XLRDError as error:
        raise WorkbookError(f"Missing domain worksheet in {path}: {error}") from error

    values: dict[str, list[str]] = defaultdict(list)
    for row_number in range(1, values_sheet.nrows):
        name = str(values_sheet.cell_value(row_number, 0)).strip()
        value = str(values_sheet.cell_value(row_number, 2))
        deleted = str(values_sheet.cell_value(row_number, 5)).strip().lower()
        enexis = str(values_sheet.cell_value(row_number, 8)).strip().lower()
        if name and value and deleted != "ja" and enexis == "x":
            values[name].append(value)

    result = {}
    for row_number in range(1, definitions.nrows):
        name = str(definitions.cell_value(row_number, 0)).strip()
        if not name:
            continue
        structured = str(definitions.cell_value(row_number, 3)).strip().lower() == "ja"
        members = tuple(sorted(values[name], key=str.casefold)) if structured else ()
        result[name] = DomainDefinition(name, structured, members)
    return result


def read_information_model(model_path: str | Path, domains_path: str | Path, version: str | None = None) -> InformationModel:
    model_file = Path(model_path)
    domain_file = Path(domains_path)
    features = _read_features(model_file)
    available_domains = _read_domains(domain_file)

    ordered_domains: list[DomainDefinition] = []
    seen_domains: set[str] = set()
    for feature in features:
        for attribute in feature.attributes:
            name = attribute.domain_name
            if not name or name in seen_domains:
                continue
            if name not in available_domains:
                raise WorkbookError(
                    f"Unknown domain {name!r} used by {feature.name}.{attribute.name} "
                    f"at Informatiemodel row {attribute.source_row}"
                )
            seen_domains.add(name)
            ordered_domains.append(available_domains[name])

    return InformationModel(version or _version_from_name(model_file), features, tuple(ordered_domains))